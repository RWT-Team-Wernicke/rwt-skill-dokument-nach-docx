#!/usr/bin/env python3
"""PDF-Extraktion nach DOCX.

Liest ein Text-PDF, erkennt seitenweise Absätze, erzeugt ein strukturiertes DOCX
mit Überschrift je Seite. Meldet Abbruch, wenn keine Textebene vorhanden ist.

Aufruf:
    python extract_pdf.py --input datei.pdf --outdir out/
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from pypdf import PdfReader
except ImportError as e:
    print(json.dumps({"fehler": f"pypdf nicht installiert: {e}"}))
    sys.exit(2)

try:
    from docx import Document
except ImportError as e:
    print(json.dumps({"fehler": f"python-docx nicht installiert: {e}"}))
    sys.exit(2)


MAX_SEITEN = 100
MAX_ZEICHEN_PRO_SEITE = 200000


def extract_pdf(input_path: Path, outdir: Path) -> dict[str, Any]:
    if not input_path.exists():
        raise FileNotFoundError(f"Eingabedatei existiert nicht: {input_path}")

    outdir.mkdir(parents=True, exist_ok=True)
    output_path = outdir / f"{input_path.stem}.docx"

    reader = PdfReader(str(input_path))
    seiten = len(reader.pages)
    if seiten > MAX_SEITEN:
        raise ValueError(
            f"PDF hat {seiten} Seiten, Höchstgrenze {MAX_SEITEN}. Bitte vorher aufteilen."
        )

    warnungen: list[str] = []
    absaetze_gesamt = 0
    leere_seiten: list[int] = []

    doc = Document()
    doc.add_paragraph(
        f"Quelle: {input_path.name} | Format: PDF | Seiten: {seiten} | Extraktion: {datetime.now(timezone.utc).isoformat(timespec='seconds')}"
    )

    for i, seite in enumerate(reader.pages, start=1):
        try:
            text = seite.extract_text() or ""
        except Exception as e:  # pypdf kann bei defekten Seiten scheitern
            warnungen.append(f"Seite {i}: Extraktion fehlgeschlagen ({e})")
            text = ""

        if len(text) > MAX_ZEICHEN_PRO_SEITE:
            warnungen.append(
                f"Seite {i}: mehr als {MAX_ZEICHEN_PRO_SEITE} Zeichen, wird gekürzt"
            )
            text = text[:MAX_ZEICHEN_PRO_SEITE]

        if not text.strip():
            leere_seiten.append(i)
            doc.add_heading(f"Seite {i}", level=1)
            doc.add_paragraph(
                "[nicht extrahierbar: keine Textebene auf dieser Seite. Wahrscheinlich Scan oder Bildinhalt.]"
            )
            continue

        doc.add_heading(f"Seite {i}", level=1)
        for absatz_roh in text.split("\n\n"):
            absatz = absatz_roh.strip()
            if absatz:
                doc.add_paragraph(absatz)
                absaetze_gesamt += 1

    if len(leere_seiten) == seiten:
        raise ValueError(
            "Keine einzige Seite enthält extrahierbaren Text. Vermutlich Scan-PDF ohne Textebene. Bitte vorher OCR-en (z. B. mit ocrmypdf) und erneut aufrufen."
        )

    if leere_seiten:
        warnungen.append(
            f"Leere Seiten (kein Text extrahierbar): {leere_seiten}"
        )

    doc.save(str(output_path))

    return {
        "eingabe": str(input_path),
        "ausgabe": str(output_path),
        "seiten": seiten,
        "leere_seiten": leere_seiten,
        "absaetze": absaetze_gesamt,
        "warnungen": warnungen,
        "status": "OK",
    }


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--outdir", required=True, type=Path)
    args = ap.parse_args()

    try:
        bericht = extract_pdf(args.input, args.outdir)
    except (FileNotFoundError, ValueError) as e:
        print(json.dumps({"fehler": str(e), "status": "FEHLER"}, ensure_ascii=False, indent=2))
        return 1
    except Exception as e:  # noqa: BLE001
        print(json.dumps({"fehler": f"Unerwarteter Fehler: {e}", "status": "FEHLER"}, ensure_ascii=False, indent=2))
        return 2

    print(json.dumps(bericht, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
