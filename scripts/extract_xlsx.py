#!/usr/bin/env python3
"""XLSX-Extraktion nach DOCX.

Liest eine Excel-Datei, extrahiert jedes Arbeitsblatt als Tabelle in ein
strukturiertes DOCX. Formeln werden als berechneter Wert übernommen.

Aufruf:
    python extract_xlsx.py --input datei.xlsx --outdir out/
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as e:
    print(json.dumps({"fehler": f"openpyxl nicht installiert: {e}"}))
    sys.exit(2)

try:
    from docx import Document
except ImportError as e:
    print(json.dumps({"fehler": f"python-docx nicht installiert: {e}"}))
    sys.exit(2)


MAX_ZELLEN_PRO_BLATT = 50000
MAX_ZEICHEN_PRO_ZELLE = 30000


def _zelltext(wert: Any) -> str:
    if wert is None:
        return ""
    if isinstance(wert, float) and wert.is_integer():
        return str(int(wert))
    return str(wert)


def extract_xlsx(input_path: Path, outdir: Path) -> dict[str, Any]:
    if not input_path.exists():
        raise FileNotFoundError(f"Eingabedatei existiert nicht: {input_path}")

    outdir.mkdir(parents=True, exist_ok=True)
    output_path = outdir / f"{input_path.stem}.docx"

    # data_only=True → berechnete Werte statt Formeltext
    wb = load_workbook(str(input_path), data_only=True, read_only=True)
    blaetter = wb.sheetnames

    warnungen: list[str] = []
    tabellen_gesamt = 0
    zellen_gesamt = 0

    doc = Document()
    doc.add_paragraph(
        f"Quelle: {input_path.name} | Format: XLSX | Blätter: {len(blaetter)} | Extraktion: {datetime.now(timezone.utc).isoformat(timespec='seconds')}"
    )

    for blattname in blaetter:
        ws = wb[blattname]
        doc.add_heading(f"Blatt: {blattname}", level=1)

        # Effektive Zellzahl bestimmen
        zellen_in_blatt = (ws.max_row or 0) * (ws.max_column or 0)
        if zellen_in_blatt == 0:
            doc.add_paragraph("[leeres Blatt]")
            continue

        if zellen_in_blatt > MAX_ZELLEN_PRO_BLATT:
            warnungen.append(
                f"Blatt '{blattname}': {zellen_in_blatt} Zellen, wird auf {MAX_ZELLEN_PRO_BLATT} beschnitten"
            )

        # Zeilen sammeln (beschnitten)
        zeilen: list[list[str]] = []
        for zeile in ws.iter_rows(values_only=True):
            gezeilt = [_zelltext(z)[:MAX_ZEICHEN_PRO_ZELLE] for z in zeile]
            if any(z for z in gezeilt):
                zeilen.append(gezeilt)
            if len(zeilen) * (ws.max_column or 1) >= MAX_ZELLEN_PRO_BLATT:
                break

        if not zeilen:
            doc.add_paragraph("[keine gefüllten Zellen]")
            continue

        # Alle Zeilen auf gleiche Spaltenzahl bringen
        max_cols = max(len(z) for z in zeilen)
        zeilen = [z + [""] * (max_cols - len(z)) for z in zeilen]

        docx_table = doc.add_table(rows=len(zeilen), cols=max_cols)
        docx_table.style = "Table Grid"
        for r, zeile in enumerate(zeilen):
            for c, wert in enumerate(zeile):
                docx_table.cell(r, c).text = wert
        tabellen_gesamt += 1
        zellen_gesamt += len(zeilen) * max_cols

    wb.close()
    doc.save(str(output_path))

    return {
        "eingabe": str(input_path),
        "ausgabe": str(output_path),
        "blaetter": blaetter,
        "tabellen": tabellen_gesamt,
        "zellen_gesamt": zellen_gesamt,
        "warnungen": warnungen,
        "status": "OK",
    }


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--outdir", required=True, type=Path)
    args = ap.parse_args()

    try:
        bericht = extract_xlsx(args.input, args.outdir)
    except (FileNotFoundError, ValueError) as e:
        print(json.dumps({"fehler": str(e), "status": "FEHLER"}, ensure_ascii=False, indent=2))
        return 1
    except Exception as e:  # noqa: BLE001
        print(json.dumps({"fehler": f"Unerwarteter Fehler: {e}", "status": "FEHLER"}, ensure_ascii=False, indent=2))
        return 2

    print(json.dumps(bericht, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
