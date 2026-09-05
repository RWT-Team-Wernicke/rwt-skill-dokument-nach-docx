#!/usr/bin/env python3
"""PPTX-Extraktion nach DOCX.

Liest eine PowerPoint-Datei, extrahiert Titel, Textrahmen, Notizen (optional),
und Tabellen jeder Folie in ein strukturiertes DOCX.

Aufruf:
    python extract_pptx.py --input datei.pptx --outdir out/ [--mit-notizen]
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError as e:
    print(json.dumps({"fehler": f"python-pptx nicht installiert: {e}"}))
    sys.exit(2)

try:
    from docx import Document
except ImportError as e:
    print(json.dumps({"fehler": f"python-docx nicht installiert: {e}"}))
    sys.exit(2)


MAX_FOLIEN = 200


def _shape_text(shape) -> str:
    """Extrahiere Text aus einer Shape, wenn vorhanden."""
    if not shape.has_text_frame:
        return ""
    lines: list[str] = []
    for p in shape.text_frame.paragraphs:
        line = "".join(r.text for r in p.runs).strip()
        if line:
            lines.append(line)
    return "\n".join(lines).strip()


def extract_pptx(input_path: Path, outdir: Path, mit_notizen: bool) -> dict[str, Any]:
    if not input_path.exists():
        raise FileNotFoundError(f"Eingabedatei existiert nicht: {input_path}")

    outdir.mkdir(parents=True, exist_ok=True)
    output_path = outdir / f"{input_path.stem}.docx"

    prs = Presentation(str(input_path))
    folien = len(prs.slides)
    if folien > MAX_FOLIEN:
        raise ValueError(
            f"Präsentation hat {folien} Folien, Höchstgrenze {MAX_FOLIEN}. Bitte vorher aufteilen."
        )

    warnungen: list[str] = []
    absaetze_gesamt = 0
    tabellen_gesamt = 0
    bilder_gesamt = 0
    notizen_gesamt = 0

    doc = Document()
    doc.add_paragraph(
        f"Quelle: {input_path.name} | Format: PPTX | Folien: {folien} | Extraktion: {datetime.now(timezone.utc).isoformat(timespec='seconds')}"
    )

    for i, folie in enumerate(prs.slides, start=1):
        # Titel bestimmen
        titel = ""
        for shape in folie.shapes:
            if shape.has_text_frame and shape == folie.shapes.title if folie.shapes.title else False:
                titel = _shape_text(shape).splitlines()[0] if _shape_text(shape) else ""
                break
        if not titel:
            # Fallback: erste Textzeile der Folie
            for shape in folie.shapes:
                t = _shape_text(shape)
                if t:
                    titel = t.splitlines()[0]
                    break
        if not titel:
            titel = f"Folie {i}"

        doc.add_heading(f"Folie {i}: {titel}", level=1)

        # Alle Shapes durchlaufen
        for shape in folie.shapes:
            # Titel wurde schon geschrieben
            if shape.has_text_frame:
                text = _shape_text(shape)
                if not text:
                    continue
                # Titel nicht doppelt schreiben
                if text.splitlines()[0] == titel and text.count("\n") == 0:
                    continue
                for line in text.splitlines():
                    if line.strip():
                        doc.add_paragraph(line.strip())
                        absaetze_gesamt += 1
            elif shape.has_table:
                pptx_table = shape.table
                zeilen = len(pptx_table.rows)
                spalten = len(pptx_table.columns)
                if zeilen == 0 or spalten == 0:
                    continue
                docx_table = doc.add_table(rows=zeilen, cols=spalten)
                docx_table.style = "Table Grid"
                for r, zeile in enumerate(pptx_table.rows):
                    for c, zelle in enumerate(zeile.cells):
                        docx_table.cell(r, c).text = zelle.text.strip()
                tabellen_gesamt += 1
            elif shape.shape_type == 13:  # Picture
                doc.add_paragraph(f"[Bild auf Folie {i}: {shape.name}]")
                bilder_gesamt += 1
            elif shape.shape_type == 6:  # Group
                warnungen.append(
                    f"Folie {i}: Gruppierte Shape ({shape.name}) wird nicht rekursiv extrahiert"
                )

        # Notizen
        if mit_notizen and folie.has_notes_slide:
            notes = folie.notes_slide.notes_text_frame.text.strip()
            if notes:
                doc.add_paragraph(f"[Notizen Folie {i}]")
                for line in notes.splitlines():
                    if line.strip():
                        doc.add_paragraph(line.strip())
                notizen_gesamt += 1

    doc.save(str(output_path))

    return {
        "eingabe": str(input_path),
        "ausgabe": str(output_path),
        "folien": folien,
        "absaetze": absaetze_gesamt,
        "tabellen": tabellen_gesamt,
        "bilder_uebersprungen": bilder_gesamt,
        "notizen_extrahiert": notizen_gesamt if mit_notizen else "deaktiviert",
        "warnungen": warnungen,
        "status": "OK",
    }


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--outdir", required=True, type=Path)
    ap.add_argument("--mit-notizen", action="store_true", help="Folien-Notizen mitextrahieren")
    args = ap.parse_args()

    try:
        bericht = extract_pptx(args.input, args.outdir, args.mit_notizen)
    except (FileNotFoundError, ValueError) as e:
        print(json.dumps({"fehler": str(e), "status": "FEHLER"}, ensure_ascii=False, indent=2))
        return 1
    except Exception as e:  # noqa: BLE001
        print(json.dumps({"fehler": f"Unerwarteter Fehler: {e}", "status": "FEHLER"}, ensure_ascii=False, indent=2))
        return 2

    print(json.dumps(bericht, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
